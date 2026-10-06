"""과제·정답 파일: 정답 수식을 실제로 계산해 기대값과 맞는지 본다."""
import io

import formulas
import pytest
from openpyxl import load_workbook

from databook import mapping as m
from databook.analysis import compute
from databook.practice import ANSWER_FILL, SKILLS, build_practice

M = 1_000_000
# formulas 라이브러리가 계산하지 못하는 함수(XLOOKUP, TOCOL, INDIRECT)가 든 과제는 값 대조에서 빼고
# 수식 문자열로 확인한다. SUBTOTAL도 계산하지 못해 그 줄만 #NAME?을 허용한다.
CALCULABLE = [s.key for s in SKILLS if s.key not in ("lookup", "transpose", "indirect")]


def evaluate(xlsx: bytes, tmp_path):
    path = tmp_path / "answer.xlsx"
    path.write_bytes(xlsx)
    solution = formulas.ExcelModel().loads(str(path)).finish().calculate()
    out = {}
    for ref, cell in solution.items():
        sheet, _, coord = ref.rpartition("!")
        if ":" not in coord and "]" in sheet and hasattr(cell, "value"):
            value = cell.value[0, 0]
            out[sheet.strip("'").split("]")[1], coord] = getattr(value, "item", lambda: value)()
    return out


def row_of(ws, label):
    return next(c.row for (c,) in ws.iter_rows(min_col=1, max_col=1) if c.value == label)


def test_every_skill_builds_both_files(db):
    practice = build_practice(db, [s.key for s in SKILLS], seed=3)
    assert len(practice.tasks) == len(SKILLS)
    task = load_workbook(io.BytesIO(practice.task_xlsx))
    answer = load_workbook(io.BytesIO(practice.answer_xlsx))
    assert task.sheetnames == answer.sheetnames
    assert task.sheetnames[0] == "안내" and {"BS", "IS", "CF", "DATA", "매크로연습"} <= set(task.sheetnames)
    assert "Sub FormatReport()" in practice.macro_code and "Sub CheckBalance()" in practice.macro_code

    yellow = ANSWER_FILL.start_color.rgb
    blanks = filled = 0
    for name in (t.sheet for t in practice.tasks):
        for row in task[name].iter_rows():
            for cell in row:
                if cell.fill.start_color.rgb == yellow:
                    assert cell.value is None, f"과제 파일 {name}!{cell.coordinate}에 답이 들어 있음"
                    blanks += 1
                    filled += answer[name][cell.coordinate].value is not None
    assert blanks > 100 and filled > 100


def test_same_seed_same_questions(db):
    a = build_practice(db, ["lookup", "index_match"], seed=1)
    b = build_practice(db, ["lookup", "index_match"], seed=1)
    names = lambda p: [c.value for (c,) in load_workbook(io.BytesIO(p.task_xlsx))["과제1"].iter_rows(min_col=1, max_col=1)]
    assert names(a) == names(b)


def test_no_skill_is_an_error(db):
    with pytest.raises(ValueError):
        build_practice(db, [])


def test_answer_formulas_compute_correctly(db, tmp_path):
    practice = build_practice(db, CALCULABLE, seed=0)
    got = evaluate(practice.answer_xlsx, tmp_path)
    wb = load_workbook(io.BytesIO(practice.answer_xlsx))
    sheet = {t.key: t.sheet for t in practice.tasks}
    subtotal_row = row_of(wb[sheet["sum"]], "③ 보이는 셀 합계 (SUBTOTAL)")
    errors = {
        k: v for k, v in got.items()
        if str(v).startswith("#") and k[1][1:] != str(subtotal_row)
    }
    assert not errors
    a = compute(db)
    cols = "BCD"  # fixture는 3개년

    def values(key, label):
        ws = wb[sheet[key]]
        r = row_of(ws, label)
        return [got.get((sheet[key].upper(), f"{c}{r}")) for c in cols]

    assert values("sum", "② 차이") == [0, 0, 0]
    r = row_of(wb[sheet["sum"]], "계정명")
    assert wb[sheet["sum"]][f"B{subtotal_row}"].value == f"=SUBTOTAL(9,B{r + 1}:B{subtotal_row - 5})"
    assert values("sumifs", m.DEBT) == a[m.DEBT]
    assert values("sumifs", "④ 순차입금 (음수면 순현금)") == a["순차입금"]
    assert values("sumifs", "③ 순운전자본") == [x + y - z for x, y, z in zip(a[m.AR], a[m.INV], a[m.AP])]
    assert values("if_error", "① 부채비율") == pytest.approx(a["부채비율"])
    assert values("if_error", "③ 영업이익률") == pytest.approx(a["영업이익률"])
    assert values("if_error", "④ 판정") == ["양호"] * 3
    assert values("growth", "① 매출액 성장률")[1:] == pytest.approx(a["매출 성장률"][1:])
    assert values("pivot", m.CASH) == a[m.CASH]

    ws = wb[sheet["growth"]]
    assert got[sheet["growth"].upper(), f"E{row_of(ws, '매출액')}"] == pytest.approx((1.0 / 0.8) ** 0.5 - 1)

    # INDEX·MATCH: 연도 순서가 뒤집힌 표에서도 계정·연도에 맞는 값
    ws = wb[sheet["index_match"]]
    head = row_of(ws, "계정명")
    by_name = {ln.name: ln for ln in db.is_}
    for r in range(head + 1, head + 6):
        name = ws[f"A{r}"].value
        if name:
            for c in cols:
                year = int(ws[f"{c}{head}"].value[2:])
                assert got[sheet["index_match"].upper(), f"{c}{r}"] == next(
                    ln for ln in db.is_ if ln.name == name).value(year)

    # 공통형 손익계산서: 매출액 대비 비율, 첫 줄은 100%
    ws = wb[sheet["abs_ref"]]
    head = row_of(ws, "계정명")
    assert [got[sheet["abs_ref"].upper(), f"{c}{head + 1}"] for c in "FGH"] == [1, 1, 1]
    assert got[sheet["abs_ref"].upper(), f"H{head + 2}"] == pytest.approx(2800 / 4000)

    # 문자 함수
    ws = wb[sheet["text"]]
    r = row_of(ws, "FY2025")
    key = sheet["text"].upper()
    assert (got[key, f"B{r}"], got[key, f"C{r}"], got[key, f"D{r}"], got[key, f"E{r}"]) == (2025, "FY", "20", "테스트전자_FY2025")
    assert got[key, f"B{row_of(ws, '⑤ IS 시트의 계정 수 (COUNTA)')}"] == len(db.is_)


def test_lookup_and_array_formulas_are_written(db):
    """계산 엔진이 지원하지 않는 함수는 수식 문자열로 확인한다."""
    practice = build_practice(db, ["lookup", "transpose"], seed=0)
    wb = load_workbook(io.BytesIO(practice.answer_xlsx))
    ws = wb["과제1"]
    r = row_of(ws, "계정명") + 1
    assert ws[f"B{r}"].value == f"=VLOOKUP($A{r},IS!$C$4:$F$14,4,FALSE)"
    assert ws[f"C{r}"].value == f"=_xlfn.XLOOKUP($A{r},IS!$C$4:$C$14,IS!$D$4:$D$14,0)"
    ws = wb["과제2"]
    r = row_of(ws, "① 연도") + 1
    assert ws[f"A{r}"].value.text == "=TRANSPOSE(IS!D3:F3)" and ws[f"A{r}"].value.ref == f"A{r}:A{r + 2}"
    assert ws[f"D{r}"].value.text == "=_xlfn.TOCOL(IS!D4:F6)" and ws[f"D{r}"].value.ref == f"D{r}:D{r + 8}"


def test_indirect_formulas(db):
    wb = load_workbook(io.BytesIO(build_practice(db, ["indirect"]).answer_xlsx))
    ws = wb["과제1"]
    top = row_of(ws, "시트 이름")
    assert ws[f"B{top}"].value == "IS"
    assert ws[f"B{top + 2}"].value == f'=INDIRECT($B${top}&"!C4")'
    assert ws[f"B{top + 3}"].value == f'=COUNTA(INDIRECT($B${top}&"!C4:C500"))'
    assert ws[f"B{top + 4}"].value == f'=INDIRECT($B${top}&"!F4")'
