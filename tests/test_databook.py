import io

import formulas
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from databook import mapping as m
from databook.analysis import CHECK_KEYS, compute
from databook.excel import AN_YEAR_COL, build_workbook
from databook.model import parse_amount

M = 1_000_000


def cls_of(db, name):
    return next(ln for ln in db.bs if ln.name == name).cls


def test_parse_amount():
    assert parse_amount("1,234") == 1234
    assert parse_amount("-50") == -50
    assert parse_amount("") is None and parse_amount("-") is None


def test_years_and_merge(db):
    assert db.years == [2023, 2024, 2025]  # 2021·2022는 데이터 없음
    assert db.is_source == "CIS"
    names = [ln.name for ln in db.bs]
    # 과거 연도에만 있는 줄은 그 해의 앞 줄 뒤에 들어간다
    assert names.index("매각예정자산") == names.index("기타유동자산") + 1


def test_bs_classification(db):
    expected = {
        "유동자산": m.SUBTOTAL, "현금및현금성자산": m.CASH, "단기금융상품": m.CASH,
        "매출채권": m.AR, "재고자산": m.INV, "기타유동금융자산": m.UNCLASSIFIED,
        "기타유동자산": m.OTHER_NWC_A, "유형자산": m.FIXED, "무형자산": m.FIXED,
        "이연법인세자산": m.OTHER_A, "매입채무": m.AP, "단기차입금": m.DEBT,
        "미지급금": m.OTHER_NWC_L, "기타유동금융부채": m.UNCLASSIFIED, "장기차입금": m.DEBT,
        "순확정급여부채": m.OTHER_L, "지배기업 소유주지분": m.SUBTOTAL, "자본금": m.EQUITY,
        "비지배지분": m.EQUITY, "자본총계": m.SUBTOTAL,
    }
    assert {name: cls_of(db, name) for name in expected} == expected


def test_tags(db):
    ni = [ln for ln in db.is_ if ln.cls == m.NI]
    assert len(ni) == 1  # 포괄손익계산서의 두 번째 당기순이익에는 태그가 없다
    assert [ln.name for ln in db.cf if ln.cls == m.DA] == ["감가상각비", "무형자산상각비"]
    assert next(ln for ln in db.is_ if ln.name == "법인세비용차감전순이익").cls == m.PBT


def test_analysis(db):
    a = compute(db)
    last = -1
    assert a["EBITDA"][last] == 700 * M
    assert a["순운전자본"][last] == (400 + 250 + 60 - 350 - 200) * M
    assert a["순차입금"][last] == (800 - 400) * M
    assert a["CAPEX"][last] == 280 * M
    assert a["매출 성장률"][0] is None
    assert a["매출 성장률"][last] == pytest.approx(1 / 9)
    for key in CHECK_KEYS:
        assert all(abs(v) <= 1 for v in a[key]), key


def test_excel_formulas_match_python(db, tmp_path):
    """엑셀 수식을 실제로 계산해 파이썬 계산값과 맞는지 본다."""
    path = tmp_path / "book.xlsx"
    path.write_bytes(build_workbook(db))

    solution = formulas.ExcelModel().loads(str(path)).finish().calculate()
    computed = {}
    for ref, cell in solution.items():
        sheet, _, coord = ref.rpartition("!")
        if ":" not in coord and hasattr(cell, "value"):
            value = cell.value[0, 0]
            computed[sheet.strip("'").split("]")[1], coord] = getattr(value, "item", lambda: value)()

    wb = load_workbook(path)

    def row_values(sheet, label):
        ws = wb[sheet]
        for (cell,) in ws.iter_rows(min_row=5, min_col=1, max_col=1):
            if cell.value == label and (sheet.upper(), f"B{cell.row}") in computed:  # 제목 줄 제외
                return [
                    computed[sheet.upper(), f"{get_column_letter(AN_YEAR_COL + i)}{cell.row}"]
                    for i in range(len(db.years))
                ]
        raise AssertionError(f"{sheet}: '{label}' 줄이 없음")

    errors = {k: v for k, v in computed.items() if isinstance(v, str) and v.startswith("#")}
    assert not errors

    a = compute(db)
    pairs = [
        ("EBITDA", "EBITDA", "EBITDA"), ("EBITDA", "매출총이익", "매출총이익"),
        ("NWC", "순운전자본 (NWC)", "순운전자본"), ("NWC", "매출채권 회전일수 (DSO)", "매출채권 회전일수"),
        ("순차입금", "순차입금", "순차입금"), ("순차입금", "순차입금 / EBITDA", "순차입금/EBITDA"),
        ("비율", "부채비율", "부채비율"), ("비율", "유동비율", "유동비율"), ("비율", "FCF", "FCF"),
        ("비율", "ROE", "ROE"), ("요약", "NWC / 매출액", "NWC/매출액"),
        ("체크", "자산총계 − 부채총계 − 자본총계", "체크: 자산-부채-자본"),
        ("체크", "자산 계정 합계 − 자산총계", "체크: 자산 분류 합계"),
        ("체크", "부채 계정 합계 − 부채총계", "체크: 부채 분류 합계"),
        ("체크", "자본 계정 합계 − 자본총계", "체크: 자본 분류 합계"),
    ]
    for sheet, label, key in pairs:
        assert row_values(sheet, label) == pytest.approx(a[key]), f"{sheet}/{label}"
    assert row_values("체크", "판정") == ["OK"] * len(db.years)
    assert row_values("요약", "검증") == ["OK"] * len(db.years)
    # 분류를 바꾸지 않았으면 '기타' 줄은 0
    assert row_values("순차입금", "기타 (엑셀에서 분류를 바꾼 계정)") == [0, 0, 0]


@pytest.mark.parametrize("fetch_name", ["fake_fetch_mixed", "fake_fetch_tree"])
def test_new_dart_format(db, fetch_name):
    """합계가 먼저 오고 코드순으로 섞인 새 형식에서도 결과가 같아야 한다."""
    import conftest
    from databook.model import build_databook

    new = build_databook(conftest.CORP, 2025, None, getattr(conftest, fetch_name))
    assert len(new.bs) == len(db.bs) - 1  # 예전 fixture에만 있는 매각예정자산 한 줄 차이
    assert {ln.name: (ln.section, ln.cls) for ln in new.bs} == {
        ln.name: (ln.section, ln.cls) for ln in db.bs if ln.name != "매각예정자산"
    }
    a, b = compute(new), compute(db)
    for key in ["순운전자본", "순차입금", "EBITDA", "CAPEX", *CHECK_KEYS]:
        assert a[key] == b[key], key

    names = [ln.name for ln in new.bs]
    assert names[0] == "유동자산" and names[-1] == "부채와자본총계"
    assert names.index("자산총계") < names.index("유동부채") < names.index("부채총계") < names.index("자본총계")
    assert [ln.name for ln in new.cf][0] == "영업활동현금흐름"
