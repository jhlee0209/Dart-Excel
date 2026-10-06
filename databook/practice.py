"""실전 과제 생성기.

고른 회사의 재무제표(BS·IS·CF 시트)에 익히고 싶은 기능별 과제 시트를 붙여 '과제 파일'을 만들고,
같은 구성에 정답 수식과 해설을 채운 '정답 파일'을 함께 만든다. 노란 칸이 답을 쓰는 자리다.
"""
from __future__ import annotations

import io
import random
from dataclasses import dataclass
from typing import Callable

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.worksheet import Worksheet

from . import mapping as m
from .model import FS_LABEL, Databook, Line

FONT = "맑은 고딕"
GREEN, GREEN_DARK, GREEN_TINT = "107C41", "0B5A2E", "E9F3EC"
ANSWER_FILL = PatternFill("solid", start_color="FFF2CC")
HEAD_FILL = PatternFill("solid", start_color=GREEN)
BRIEF_FILL = PatternFill("solid", start_color=GREEN_TINT)
GREY, RED = "7F7F7F", "C00000"

MIL = '#,##0,,;(#,##0,,);"-"'  # 셀 값은 원, 표시는 백만원
PCT = '0.0%;-0.0%;"-"'
INT = "#,##0"

HDR, FIRST = 3, 4  # 재무제표 시트의 머리글 행, 첫 데이터 행
YEAR_COL = 4  # 재무제표 시트의 첫 연도 열(D)

G_DAILY, G_SOMETIMES, G_EXTRA = "함수 · 매일 쓰는 기본", "함수 · 가끔 쓰는 응용", "함수 · 함께 알아두기"
G_FEATURE, G_SHORTCUT, G_MACRO = "엑셀 기능", "단축키", "매크로"


def _font(bold=False, color="000000", size=10, italic=False, name=FONT) -> Font:
    return Font(name=name, size=size, bold=bold, color=color, italic=italic)


@dataclass
class Task:
    key: str
    sheet: str
    title: str
    functions: str
    brief: list[str]


@dataclass
class Practice:
    task_xlsx: bytes
    answer_xlsx: bytes
    tasks: list[Task]
    macro_code: str | None  # 매크로 정답 코드(붙여넣기용 텍스트)


class Ctx:
    """과제 하나를 만들 때 필요한 것: 회사 데이터, 시트 위치, 정답 여부."""

    def __init__(self, db: Databook, wb: Workbook, answer: bool, seed: int):
        self.db, self.wb, self.answer = db, wb, answer
        self.years = db.years
        self.n = len(db.years)
        self.rand = random.Random(seed)
        self.lines = {"BS": db.bs, "IS": db.is_, "CF": db.cf}
        self.last = {name: FIRST + max(len(lines), 1) - 1 for name, lines in self.lines.items()}
        self._row = {id(ln): FIRST + i for name in self.lines for i, ln in enumerate(self.lines[name])}

    # 재무제표 시트 주소
    def ycol(self, i: int) -> str:
        return L(YEAR_COL + i)

    def col_range(self, sheet: str, col: str, lock_col: bool = True) -> str:
        c = f"${col}" if lock_col else col
        return f"{sheet}!{c}${FIRST}:{c}${self.last[sheet]}"

    def cell(self, sheet: str, line: Line, i: int) -> str:
        return f"{sheet}!{self.ycol(i)}{self._row[id(line)]}"

    def row(self, line: Line) -> int:
        return self._row[id(line)]

    # 줄 찾기
    def tagged(self, sheet: str, tag: str) -> Line | None:
        return next((ln for ln in self.lines[sheet] if ln.cls == tag), None)

    def role(self, role: str) -> Line | None:
        return self.db.role_line(role)

    def is_picks(self, k: int) -> list[Line]:
        """손익계산서에서 이름이 겹치지 않고 최근 연도 값이 있는 줄을 k개까지 고른다(순서 유지)."""
        seen, pool = set(), []
        for ln in self.db.is_:
            if ln.value(self.years[-1]) and ln.name not in seen and abs(ln.value(self.years[-1])) > 1_000_000:
                pool.append(ln)
            seen.add(ln.name)
        chosen = set(map(id, self.rand.sample(pool, min(k, len(pool)))))
        return [ln for ln in pool if id(ln) in chosen]


def display_name(line: Line) -> str:
    """합계 줄은 회사마다 표기가 달라(자 산 총 계 등) 표준 이름으로 통일한다."""
    return line.role or line.name


class Page:
    """과제 시트 한 장."""

    def __init__(self, ctx: Ctx, name: str, title: str, brief: list[str]):
        self.ctx = ctx
        self.ws: Worksheet = ctx.wb.create_sheet(name)
        self.ws.sheet_view.showGridLines = False
        self.ws["A1"] = title + ("  — 정답" if ctx.answer else "")
        self.ws["A1"].font = _font(bold=True, size=14, color=GREEN_DARK)
        for i, text in enumerate(brief):
            for col in range(1, 10):
                self.ws.cell(row=3 + i, column=col).fill = BRIEF_FILL
            self.ws.cell(row=3 + i, column=1, value=text).font = _font()
        self.top = 3 + len(brief) + 1  # 표를 시작할 행
        self.ws.column_dimensions["A"].width = 34
        for col in range(2, 12):
            self.ws.column_dimensions[L(col)].width = 15

    def head(self, row: int, labels: list[str], col: int = 1) -> None:
        for i, label in enumerate(labels):
            c = self.ws.cell(row=row, column=col + i, value=label)
            c.font, c.fill = _font(bold=True, color="FFFFFF"), HEAD_FILL
            c.alignment = Alignment(horizontal="left" if col + i == 1 else "right")

    def put(self, ref: str, value, fmt: str | None = None, bold: bool = False) -> None:
        c = self.ws[ref]
        c.value, c.font = value, _font(bold=bold)
        if fmt:
            c.number_format = fmt

    def ans(self, ref: str, formula: str, fmt: str | None = None) -> None:
        """답을 쓰는 노란 칸. 정답 파일에서만 수식이 들어간다."""
        c = self.ws[ref]
        c.fill = ANSWER_FILL
        c.font = _font(color=GREEN_DARK)
        if fmt:
            c.number_format = fmt
        if self.ctx.answer:
            c.value = "=" + formula

    def ans_array(self, ref_range: str, formula: str, fmt: str | None = None) -> None:
        """여러 칸에 걸친 배열 수식 답."""
        for row in self.ws[ref_range]:
            for c in row:
                c.fill, c.font = ANSWER_FILL, _font(color=GREEN_DARK)
                if fmt:
                    c.number_format = fmt
        if self.ctx.answer:
            first = ref_range.split(":")[0]
            self.ws[first] = ArrayFormula(ref_range, "=" + formula)

    def note(self, row: int, text: str, col: int = 1) -> None:
        """해설. 정답 파일에만 나온다."""
        if self.ctx.answer:
            self.ws.cell(row=row, column=col, value="해설  " + text).font = _font(color=GREY, italic=True)

    def notes(self, row: int, texts: list[str]) -> int:
        for i, text in enumerate(texts):
            self.note(row + i, text)
        return row + len(texts)


# ── 재무제표·안내 시트 ───────────────────────────────────────────────────
def _statement_sheet(ctx: Ctx, name: str, title: str, tag_label: str) -> None:
    ws = ctx.wb.create_sheet(name)
    ws["A1"] = f"{title} — {ctx.db.corp['corp_name']} ({FS_LABEL[ctx.db.fs_div]})"
    ws["A1"].font = _font(bold=True, size=13, color=GREEN_DARK)
    ws["A2"] = "셀 값은 원 단위이고, 표시형식으로 백만원을 보여 줍니다. 출처: 금융감독원 DART"
    ws["A2"].font = _font(color=GREY, italic=True)
    labels = ["구분", tag_label, "계정명"] + [f"FY{y}" for y in ctx.years]
    for i, label in enumerate(labels, 1):
        c = ws.cell(row=HDR, column=i, value=label)
        c.font, c.fill = _font(bold=True, color="FFFFFF"), HEAD_FILL
        c.alignment = Alignment(horizontal="left" if i <= 3 else "right")
    for line in ctx.lines[name]:
        r = ctx.row(line)
        total = name == "BS" and line.cls == m.SUBTOTAL
        ws.cell(row=r, column=1, value=line.section or None).font = _font(color=GREY)
        ws.cell(row=r, column=2, value=line.cls or None).font = _font()
        ws.cell(row=r, column=3, value=display_name(line)).font = _font(bold=total)
        for i, year in enumerate(ctx.years):
            c = ws.cell(row=r, column=YEAR_COL + i, value=line.values.get(year))
            c.font, c.number_format = _font(bold=total), MIL
    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 17
    ws.column_dimensions["C"].width = 40
    for i in range(ctx.n):
        ws.column_dimensions[ctx.ycol(i)].width = 15
    ws.freeze_panes = ws.cell(row=FIRST, column=YEAR_COL)


def _guide_sheet(ctx: Ctx, tasks: list[Task]) -> None:
    ws = ctx.wb.create_sheet("안내", 0)
    ws.sheet_view.showGridLines = False
    corp = ctx.db.corp["corp_name"]
    ws["A1"] = f"{corp} 실전 과제" + (" — 정답" if ctx.answer else "")
    ws["A1"].font = _font(bold=True, size=16, color=GREEN_DARK)
    rules = [
        f"자료: {corp} {FS_LABEL[ctx.db.fs_div]}재무제표 FY{ctx.years[0]}~FY{ctx.years[-1]} (BS · IS · CF 시트)",
        "노란 칸이 답을 쓰는 자리입니다. 숫자를 직접 입력하지 말고 수식으로 채우세요.",
        "한 칸에 수식을 쓰고 옆·아래로 복사해서 표를 완성하는 것이 목표입니다($ 고정에 주의).",
        "시간 제한은 없습니다. 막히면 '기능 학습'을 다시 보고, 다 풀고 나서 정답 파일과 비교하세요.",
        "정답과 수식이 달라도 값이 같고, 복사해도 깨지지 않으면 좋은 답입니다.",
    ]
    if ctx.answer:
        rules = ["노란 칸에 정답 수식이 들어 있고, 회색 글씨가 해설입니다."] + rules[:1]
    for i, text in enumerate(rules):
        ws.cell(row=3 + i, column=1, value="· " + text).font = _font()
    top = 4 + len(rules)
    for i, label in enumerate(["번호", "시트", "과제", "쓰는 기능"], 1):
        c = ws.cell(row=top, column=i, value=label)
        c.font, c.fill = _font(bold=True, color="FFFFFF"), HEAD_FILL
    for n, task in enumerate(tasks, 1):
        for i, value in enumerate([n, task.sheet, task.title, task.functions], 1):
            ws.cell(row=top + n, column=i, value=value).font = _font()
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 44
    ws.column_dimensions["D"].width = 44


# ── 과제 ────────────────────────────────────────────────────────────────
def _year_heads(ctx: Ctx, first: str) -> list[str]:
    return [first] + [f"FY{y}" for y in ctx.years]


def task_sum(ctx: Ctx, sheet: str) -> Task:
    section = m.CUR_A if ctx.role(m.CUR_A) else m.SEC_A
    total = ctx.role(m.CUR_A) or ctx.role("자산총계")
    lines = [ln for ln in ctx.db.bs if ln.section == section and ln.cls != m.SUBTOTAL]
    brief = [
        f"재무상태표의 {section} 세부 계정입니다. 노란 칸을 수식으로 채우세요.",
        "① 합계: SUM   ② 차이: 합계 − 재무상태표의 값   ③ 보이는 셀 합계: SUBTOTAL(9, …)",
        "④ 다 채운 뒤 표에 필터를 걸고 계정 몇 개를 숨겨 보세요. ①과 ③이 어떻게 달라지나요?",
    ]
    p = Page(ctx, sheet, "과제: SUM · SUBTOTAL", brief)
    top = p.top
    p.head(top, _year_heads(ctx, "계정명"))
    for k, ln in enumerate(lines):
        r = top + 1 + k
        p.put(f"A{r}", display_name(ln))
        for i in range(ctx.n):
            p.put(f"{L(2 + i)}{r}", f"={ctx.cell('BS', ln, i)}", MIL)
    a, b = top + 1, top + len(lines)
    r = b + 2
    total_name = display_name(total) if total else section
    for label in ("① 합계 (SUM)", f"   재무상태표의 {total_name}", "② 차이", "③ 보이는 셀 합계 (SUBTOTAL)"):
        p.put(f"A{r}", label, bold=not label.startswith(" "))
        r += 1
    for i in range(ctx.n):
        c = L(2 + i)
        p.ans(f"{c}{b + 2}", f"SUM({c}{a}:{c}{b})", MIL)
        p.put(f"{c}{b + 3}", f"={ctx.cell('BS', total, i)}" if total else 0, MIL)
        p.ans(f"{c}{b + 4}", f"{c}{b + 2}-{c}{b + 3}", MIL)
        p.ans(f"{c}{b + 5}", f"SUBTOTAL(9,{c}{a}:{c}{b})", MIL)
    p.notes(b + 7, [
        "차이가 0이면 세부 계정이 빠짐없이 들어왔다는 뜻입니다. 자료를 받으면 가장 먼저 하는 확인입니다.",
        "필터로 행을 숨기면 SUM은 그대로이고 SUBTOTAL만 줄어듭니다. SUBTOTAL은 보이는 셀만 계산합니다.",
    ])
    return Task("sum", sheet, f"{section} 합계와 검증", "SUM, SUBTOTAL", brief)


def task_sumifs(ctx: Ctx, sheet: str) -> Task:
    present = {ln.cls for ln in ctx.db.bs}
    classes = [c for c in (m.AR, m.INV, m.AP, m.DEBT, m.CASH) if c in present]
    brief = [
        "BS 시트 B열에는 계정별 '분류'가 있습니다. 분류별 합계를 SUMIFS로 구하세요.",
        "① 분류 × 연도 합계: 한 칸에 쓰고 오른쪽·아래로 복사   ② 계정 수: COUNTIFS",
        "③ 순운전자본 = 매출채권 + 재고자산 − 매입채무   ④ 순차입금 = 차입금 − 현금성자산",
    ]
    p = Page(ctx, sheet, "과제: SUMIFS · COUNTIFS", brief)
    top = p.top
    count_col = L(2 + ctx.n)
    p.head(top, _year_heads(ctx, "분류") + ["계정 수"])
    rows = {}
    for k, cls in enumerate(classes):
        r = rows[cls] = top + 1 + k
        p.put(f"A{r}", cls)
        for i in range(ctx.n):
            p.ans(f"{L(2 + i)}{r}",
                  f"SUMIFS({ctx.col_range('BS', ctx.ycol(i), lock_col=False)},{ctx.col_range('BS', 'B')},$A{r})", MIL)
        p.ans(f"{count_col}{r}", f"COUNTIFS({ctx.col_range('BS', 'B')},$A{r})", INT)

    def combo(i: int, plus: list[str], minus: list[str]) -> str:
        c = L(2 + i)
        terms = [f"{c}{rows[x]}" for x in plus if x in rows]
        body = "+".join(terms) or "0"
        return body + "".join(f"-{c}{rows[x]}" for x in minus if x in rows)

    r = top + len(classes) + 2
    p.put(f"A{r}", "③ 순운전자본", bold=True)
    p.put(f"A{r + 1}", "④ 순차입금 (음수면 순현금)", bold=True)
    for i in range(ctx.n):
        p.ans(f"{L(2 + i)}{r}", combo(i, [m.AR, m.INV], [m.AP]), MIL)
        p.ans(f"{L(2 + i)}{r + 1}", combo(i, [m.DEBT], [m.CASH]), MIL)
    p.notes(r + 3, [
        "합할 범위는 행만 고정(D$4:D$n)해서 오른쪽으로 복사하면 다음 연도로 넘어가게 하고, 조건 범위(B열)는 모두 고정합니다.",
        "조건은 \"차입금\"처럼 글자를 직접 쓰지 않고 $A열을 참조합니다. 분류 이름이 바뀌어도 수식을 고칠 필요가 없습니다.",
    ])
    return Task("sumifs", sheet, "분류별 집계로 순운전자본·순차입금 구하기", "SUMIFS, COUNTIFS", brief)


def task_lookup(ctx: Ctx, sheet: str) -> Task:
    picks = ctx.is_picks(5)
    first, last = ctx.years[0], ctx.years[-1]
    brief = [
        "아래 계정의 금액을 IS 시트에서 찾아오세요.",
        f"① FY{last}: VLOOKUP (마지막 인수는 FALSE)   ② FY{first}: XLOOKUP   ③ 증감 = ① − ②",
        "XLOOKUP은 Microsoft 365·Excel 2021 이상에서만 됩니다.",
    ]
    p = Page(ctx, sheet, "과제: VLOOKUP · XLOOKUP", brief)
    top = p.top
    p.head(top, ["계정명", f"① FY{last} (VLOOKUP)", f"② FY{first} (XLOOKUP)", "③ 증감"])
    table = f"IS!$C${FIRST}:${ctx.ycol(ctx.n - 1)}${ctx.last['IS']}"
    for k, ln in enumerate(picks):
        r = top + 1 + k
        p.put(f"A{r}", display_name(ln))
        p.ans(f"B{r}", f"VLOOKUP($A{r},{table},{ctx.n + 1},FALSE)", MIL)
        p.ans(f"C{r}", f"_xlfn.XLOOKUP($A{r},{ctx.col_range('IS', 'C')},{ctx.col_range('IS', 'D')},0)", MIL)
        p.ans(f"D{r}", f"B{r}-C{r}", MIL)
    for col in "BCD":
        p.ws.column_dimensions[col].width = 24
    p.notes(top + len(picks) + 2, [
        f"VLOOKUP의 열 번호 {ctx.n + 1}은 표 범위(C열부터)에서 FY{last}가 몇 번째 열인지입니다. 열이 삽입되면 틀어지는 것이 약점입니다.",
        "XLOOKUP은 찾을 범위와 가져올 범위를 따로 주므로 열 번호가 필요 없고, 마지막 인수로 못 찾았을 때 값을 정할 수 있습니다.",
        "같은 이름의 계정이 여러 줄이면 두 함수 모두 첫 번째 줄을 가져옵니다. 찾기 전에 이름이 유일한지 COUNTIFS로 확인하는 습관을 들이세요.",
    ])
    return Task("lookup", sheet, "계정명으로 금액 찾아오기", "VLOOKUP, XLOOKUP", brief)


def task_text(ctx: Ctx, sheet: str) -> Task:
    brief = [
        "A열의 연도 표기에서 필요한 부분을 잘라내고, 찾기용 키를 만드세요.",
        "① 연도 숫자: RIGHT로 4글자 (숫자로 바꾸려면 VALUE)   ② 앞 두 글자: LEFT   ③ 가운데 두 글자(20): MID",
        "④ 키: CONCATENATE로 '회사명_FY2025' 형태   ⑤ IS 시트의 계정이 몇 줄인지: COUNTA",
    ]
    p = Page(ctx, sheet, "과제: LEFT · RIGHT · MID · CONCATENATE · COUNTA", brief)
    top = p.top
    p.put(f"A{top}", "회사명", bold=True)
    p.put(f"B{top}", ctx.db.corp["corp_name"])
    h = top + 2
    p.head(h, ["연도 표기", "① 연도 (RIGHT)", "② 앞 두 글자 (LEFT)", "③ 가운데 (MID)", "④ 키 (CONCATENATE)"])
    for k, year in enumerate(ctx.years):
        r = h + 1 + k
        p.put(f"A{r}", f"FY{year}")
        p.ans(f"B{r}", f"VALUE(RIGHT(A{r},4))", "0")
        p.ans(f"C{r}", f"LEFT(A{r},2)")
        p.ans(f"D{r}", f"MID(A{r},3,2)")
        p.ans(f"E{r}", f'CONCATENATE($B${top},"_",A{r})')
    r = h + ctx.n + 2
    p.put(f"A{r}", "⑤ IS 시트의 계정 수 (COUNTA)", bold=True)
    p.ans(f"B{r}", f"COUNTA({ctx.col_range('IS', 'C')})", INT)
    for col in "BCDE":
        p.ws.column_dimensions[col].width = 24
    p.notes(r + 2, [
        "RIGHT·LEFT·MID의 결과는 문자입니다. ①을 VALUE로 감싸지 않으면 숫자 2025와 비교했을 때 다르다고 나옵니다.",
        "④처럼 두 값을 이어 붙인 키 열을 만들어 두면, 조건이 두 개인 찾기도 VLOOKUP 하나로 할 수 있습니다.",
        "COUNTA로 받은 자료의 행 수를 원본과 맞춰 보는 것도 자주 쓰는 검증입니다.",
    ])
    return Task("text", sheet, "문자 자르기와 이어 붙이기", "LEFT, RIGHT, MID, CONCATENATE, COUNTA", brief)


def task_index_match(ctx: Ctx, sheet: str) -> Task:
    picks = ctx.is_picks(5)
    order = list(reversed(range(ctx.n)))  # 연도 순서를 뒤집어, 열 위치도 MATCH로 찾게 한다
    brief = [
        "계정(세로)과 연도(가로)를 동시에 찾아 표를 채우세요. 연도 순서가 IS 시트와 반대입니다.",
        "INDEX(값 범위, MATCH(계정명, 계정명 범위, 0), MATCH(연도, 연도 머리글 범위, 0))",
        "수식 하나를 써서 표 전체에 복사할 수 있어야 합니다.",
    ]
    p = Page(ctx, sheet, "과제: INDEX · MATCH", brief)
    top = p.top
    p.head(top, ["계정명"] + [f"FY{ctx.years[i]}" for i in order])
    lastc = ctx.ycol(ctx.n - 1)
    values = f"IS!$D${FIRST}:${lastc}${ctx.last['IS']}"
    heads = f"IS!$D${HDR}:${lastc}${HDR}"
    for k, ln in enumerate(picks):
        r = top + 1 + k
        p.put(f"A{r}", display_name(ln))
        for j in range(ctx.n):
            c = L(2 + j)
            p.ans(f"{c}{r}", f"INDEX({values},MATCH($A{r},{ctx.col_range('IS', 'C')},0),MATCH({c}${top},{heads},0))", MIL)
    p.notes(top + len(picks) + 2, [
        "$A열은 열만, 머리글 행은 행만 고정합니다(혼합 참조). 그래야 복사했을 때 계정은 아래로, 연도는 옆으로 따라갑니다.",
        "MATCH의 세 번째 인수 0은 정확히 일치입니다. 빠뜨리면 정렬되지 않은 표에서 엉뚱한 값을 가져옵니다.",
        "VLOOKUP과 달리 열 번호를 손으로 세지 않으므로, 원본 표의 열 순서가 바뀌어도 맞게 가져옵니다.",
    ])
    return Task("index_match", sheet, "계정 × 연도 양방향 찾기", "INDEX, MATCH", brief)


def task_transpose(ctx: Ctx, sheet: str) -> Task:
    line = ctx.tagged("IS", m.REV) or ctx.db.is_[0]
    lastc = ctx.ycol(ctx.n - 1)
    rows3 = min(3, len(ctx.db.is_))
    brief = [
        f"IS 시트에 가로로 놓인 '{display_name(line)}'를 세로로 돌려 가져오세요.",
        "① 연도 머리글과 금액을 각각 TRANSPOSE로 (구버전은 범위를 선택하고 Ctrl + Shift + Enter)",
        f"② IS 시트 위 {rows3}줄의 금액 전체를 TOCOL로 세로 한 줄로 펼치기 (Microsoft 365 전용)",
    ]
    p = Page(ctx, sheet, "과제: TRANSPOSE · TOCOL", brief)
    top = p.top
    p.head(top, ["① 연도", display_name(line)])
    a, b = top + 1, top + ctx.n
    p.ans_array(f"A{a}:A{b}", f"TRANSPOSE(IS!D{HDR}:{lastc}{HDR})")
    p.ans_array(f"B{a}:B{b}", f"TRANSPOSE(IS!D{ctx.row(line)}:{lastc}{ctx.row(line)})", MIL)
    p.head(top, ["② 펼친 금액 (TOCOL)"], col=4)
    p.ws.column_dimensions["D"].width = 24
    p.ans_array(f"D{a}:D{a + rows3 * ctx.n - 1}", f"_xlfn.TOCOL(IS!D{FIRST}:{lastc}{FIRST + rows3 - 1})", MIL)
    p.notes(max(b, a + rows3 * ctx.n - 1) + 2, [
        "값으로 한 번만 돌리면 될 때는 복사 후 '행열 바꿔 붙여넣기'가 더 빠릅니다. 원본을 따라 바뀌어야 할 때 함수를 씁니다.",
        "TOCOL은 넓은 표를 '계정-연도-금액'의 긴 형태로 바꿔 피벗 테이블에 넣을 때 씁니다. 구버전에서는 #NAME? 오류가 납니다.",
    ])
    return Task("transpose", sheet, "가로 표를 세로로 바꾸기", "TRANSPOSE, TOCOL", brief)


def task_indirect(ctx: Ctx, sheet: str) -> Task:
    brief = [
        "B열의 시트 이름을 바꾸면 결과가 따라 바뀌도록 INDIRECT로 수식을 쓰세요. (BS / IS / CF 중에서 선택)",
        "① 그 시트의 첫 번째 계정명 (C4 셀)   ② 그 시트의 계정 수: COUNTA와 함께",
        f"③ 그 시트 첫 줄의 FY{ctx.years[-1]} 금액",
    ]
    p = Page(ctx, sheet, "과제: INDIRECT", brief)
    top = p.top
    p.put(f"A{top}", "시트 이름", bold=True)
    p.put(f"B{top}", "IS")
    p.ws[f"B{top}"].font = _font(color="0000FF")
    validation = DataValidation(type="list", formula1='"BS,IS,CF"')
    p.ws.add_data_validation(validation)
    validation.add(f"B{top}")
    lastc = ctx.ycol(ctx.n - 1)
    p.put(f"A{top + 2}", "① 첫 번째 계정명", bold=True)
    p.ans(f"B{top + 2}", f'INDIRECT($B${top}&"!C{FIRST}")')
    p.put(f"A{top + 3}", "② 계정 수", bold=True)
    p.ans(f"B{top + 3}", f'COUNTA(INDIRECT($B${top}&"!C{FIRST}:C500"))', INT)
    p.put(f"A{top + 4}", f"③ 첫 줄의 FY{ctx.years[-1]} 금액", bold=True)
    p.ans(f"B{top + 4}", f'INDIRECT($B${top}&"!{lastc}{FIRST}")', MIL)
    p.ws.column_dimensions["B"].width = 30
    p.notes(top + 6, [
        "INDIRECT는 \"IS!C4\" 같은 문자를 실제 주소로 바꿉니다. 시트 이름에 공백이 있으면 작은따옴표로 감싸야 합니다: \"'\"&B5&\"'!C4\"",
        "편리하지만 참조 추적이 안 되고 파일이 느려집니다. 시트가 수십 개로 같은 구조일 때처럼 꼭 필요한 경우에만 씁니다.",
    ])
    return Task("indirect", sheet, "시트 이름으로 참조 바꾸기", "INDIRECT, COUNTA", brief)


def task_abs_ref(ctx: Ctx, sheet: str) -> Task:
    base = ctx.tagged("IS", m.REV) or ctx.db.is_[0]
    start = ctx.db.is_.index(base)
    lines = [ln for ln in ctx.db.is_[start:start + 12] if any(ln.value(y) for y in ctx.years)]
    brief = [
        f"왼쪽 금액을 '{display_name(base)}' 대비 비율로 바꿔 오른쪽 표(공통형 손익계산서)를 채우세요.",
        "수식을 맨 왼쪽 위 한 칸에만 쓰고, 오른쪽과 아래로 복사해서 완성해야 합니다.",
        "힌트: 분모는 행만 고정합니다 (F4를 두 번).",
    ]
    p = Page(ctx, sheet, "과제: 절대·혼합 참조", brief)
    top = p.top
    gap = 2 + ctx.n  # 금액 표와 비율 표 사이 빈 열
    p.head(top, _year_heads(ctx, "계정명"))
    p.head(top, [f"FY{y} 비율" for y in ctx.years], col=gap + 1)
    p.ws.column_dimensions[L(gap)].width = 3
    for k, ln in enumerate(lines):
        r = top + 1 + k
        p.put(f"A{r}", display_name(ln))
        for i in range(ctx.n):
            p.put(f"{L(2 + i)}{r}", f"={ctx.cell('IS', ln, i)}", MIL)
            p.ans(f"{L(gap + 1 + i)}{r}", f"{L(2 + i)}{r}/{L(2 + i)}${top + 1}", PCT)
    p.notes(top + len(lines) + 2, [
        f"분모 B${top + 1}은 행만 고정했습니다. 아래로 복사해도 매출액 행을 가리키고, 오른쪽으로 복사하면 다음 연도로 넘어갑니다.",
        "$B$5처럼 둘 다 고정하면 오른쪽으로 복사했을 때 모든 연도가 첫 해 매출로 나뉩니다. 가장 흔한 실수입니다.",
    ])
    return Task("abs_ref", sheet, "공통형 손익계산서 만들기", "$ 절대·혼합 참조", brief)


def task_if_error(ctx: Ctx, sheet: str) -> Task:
    brief = [
        "아래 금액으로 비율을 계산하고 판정을 내리세요. 분모가 0이어도 오류가 보이지 않아야 합니다.",
        '① 부채비율 = 부채총계 ÷ 자본총계   ② 유동비율 = 유동자산 ÷ 유동부채   ③ 영업이익률  (모두 IFERROR로 감싸고 오류면 "n/a")',
        '④ 판정: 부채비율이 200% 이하이고 유동비율이 100% 이상이면 "양호", 아니면 "주의" (IF와 AND)',
    ]
    p = Page(ctx, sheet, "과제: IF · AND · IFERROR", brief)
    top = p.top
    p.head(top, _year_heads(ctx, "항목"))
    given = [
        ("부채총계", "BS", ctx.role("부채총계")), ("자본총계", "BS", ctx.role("자본총계")),
        ("유동자산", "BS", ctx.role(m.CUR_A)), ("유동부채", "BS", ctx.role(m.CUR_L)),
        ("매출액", "IS", ctx.tagged("IS", m.REV)), ("영업이익", "IS", ctx.tagged("IS", m.EBIT)),
    ]
    rows = {}
    for k, (label, sh, ln) in enumerate(given):
        r = rows[label] = top + 1 + k
        p.put(f"A{r}", label)
        for i in range(ctx.n):
            p.put(f"{L(2 + i)}{r}", f"={ctx.cell(sh, ln, i)}" if ln else 0, MIL)
    r0 = top + len(given) + 2
    ratios = [("① 부채비율", "부채총계", "자본총계"), ("② 유동비율", "유동자산", "유동부채"), ("③ 영업이익률", "영업이익", "매출액")]
    for k, (label, num, den) in enumerate(ratios):
        p.put(f"A{r0 + k}", label, bold=True)
        for i in range(ctx.n):
            c = L(2 + i)
            p.ans(f"{c}{r0 + k}", f'IFERROR({c}{rows[num]}/{c}{rows[den]},"n/a")', PCT)
    p.put(f"A{r0 + 3}", "④ 판정", bold=True)
    for i in range(ctx.n):
        c = L(2 + i)
        p.ans(f"{c}{r0 + 3}", f'IF(AND({c}{r0}<=2,{c}{r0 + 1}>=1),"양호","주의")')
        p.ws[f"{c}{r0 + 3}"].alignment = Alignment(horizontal="right")
    p.notes(r0 + 5, [
        "200%는 수식에서 2, 100%는 1입니다. 퍼센트는 표시형식일 뿐 값은 소수입니다.",
        "IFERROR는 오류를 가립니다. 나눗셈처럼 원인이 분명한 곳에만 쓰고, 수식 전체를 습관적으로 감싸지는 않습니다.",
        "기준값(200%, 100%)을 별도 셀에 두고 참조하면 기준이 바뀔 때 한 곳만 고치면 됩니다. 한 번 그렇게 바꿔 보세요.",
    ])
    return Task("if_error", sheet, "재무비율 계산과 판정", "IF, AND, IFERROR", brief)


def task_growth(ctx: Ctx, sheet: str) -> Task:
    brief = [
        "매출액과 영업이익의 성장률을 구하세요.",
        "① 전년 대비 성장률 = 당기 ÷ 전기 − 1   ② 연평균 성장률(CAGR) = (마지막 ÷ 처음) ^ (1 ÷ 기간) − 1",
        f"기간은 연도 수가 아니라 간격입니다: FY{ctx.years[0]}~FY{ctx.years[-1]}는 {ctx.n - 1}년.",
    ]
    p = Page(ctx, sheet, "과제: 성장률 · CAGR", brief)
    top = p.top
    p.head(top, _year_heads(ctx, "항목") + ["② CAGR"])
    items = [("매출액", ctx.tagged("IS", m.REV)), ("영업이익", ctx.tagged("IS", m.EBIT))]
    cagr_col, firstc, lastc = L(2 + ctx.n), "B", L(1 + ctx.n)
    for k, (label, ln) in enumerate(items):
        r, g = top + 1 + k, top + 4 + k
        p.put(f"A{r}", label)
        p.put(f"A{g}", f"① {label} 성장률", bold=True)
        for i in range(ctx.n):
            c = L(2 + i)
            p.put(f"{c}{r}", f"={ctx.cell('IS', ln, i)}" if ln else 0, MIL)
            if i:
                p.ans(f"{c}{g}", f'IFERROR({c}{r}/{L(1 + i)}{r}-1,"n/a")', PCT)
        p.ans(f"{cagr_col}{r}", f'IFERROR(({lastc}{r}/{firstc}{r})^(1/{ctx.n - 1})-1,"n/a")', PCT)
    p.notes(top + 7, [
        "처음 값이 음수이거나 0이면 CAGR은 의미가 없어 오류가 납니다. 영업이익이 적자였던 회사에서 확인해 보세요.",
        "=RRI(기간, 처음, 마지막) 함수로도 같은 값을 구할 수 있습니다.",
    ])
    return Task("growth", sheet, "전년 대비 성장률과 CAGR", "사칙연산, ^, IFERROR", brief)


def task_cond_format(ctx: Ctx, sheet: str) -> Task:
    rev, ebit = ctx.tagged("IS", m.REV), ctx.tagged("IS", m.EBIT)
    brief = [
        "아래 표에 조건부 서식을 걸어, 눈으로 훑기만 해도 이상한 해가 보이게 만드세요. (홈 › 조건부 서식)",
        "① 매출 성장률이 음수인 칸: 빨간 글씨 + 연한 빨강 채우기   (셀 강조 규칙 › 보다 작음 › 0)",
        "② 영업이익률 줄: 3색 색조 (낮음 빨강 → 중간 흰색 → 높음 초록)",
    ]
    p = Page(ctx, sheet, "과제: 조건부 서식", brief)
    top = p.top
    p.head(top, _year_heads(ctx, "항목"))
    p.put(f"A{top + 1}", "매출 성장률")
    p.put(f"A{top + 2}", "영업이익률")
    for i in range(ctx.n):
        c = L(2 + i)
        if i and rev:
            p.put(f"{c}{top + 1}", f'=IFERROR({ctx.cell("IS", rev, i)}/{ctx.cell("IS", rev, i - 1)}-1,"n/a")', PCT)
        if rev and ebit:
            p.put(f"{c}{top + 2}", f'=IFERROR({ctx.cell("IS", ebit, i)}/{ctx.cell("IS", rev, i)},"n/a")', PCT)
    if ctx.answer:
        last = L(1 + ctx.n)
        p.ws.conditional_formatting.add(
            f"B{top + 1}:{last}{top + 1}",
            CellIsRule(operator="lessThan", formula=["0"], font=Font(color=RED),
                       fill=PatternFill("solid", start_color="F8CBAD", end_color="F8CBAD")),
        )
        p.ws.conditional_formatting.add(
            f"B{top + 2}:{last}{top + 2}",
            ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50,
                           mid_color="FFFFFF", end_type="max", end_color="63BE7B"),
        )
    p.notes(top + 4, [
        "이 시트에는 두 규칙이 이미 걸려 있습니다. 홈 › 조건부 서식 › 규칙 관리에서 확인하세요.",
        "조건부 서식은 값이 바뀌면 색도 따라 바뀝니다. 손으로 칠한 색과 달리 다음 분기 자료를 넣어도 그대로 작동합니다.",
        "규칙이 많아지면 파일이 느려지고 관리가 어렵습니다. 범위를 필요한 만큼만 잡으세요.",
    ])
    return Task("cond_format", sheet, "이상한 해가 한눈에 보이게 하기", "조건부 서식", brief)


def task_pivot(ctx: Ctx, sheet: str) -> Task:
    data = ctx.wb.create_sheet("DATA")
    for i, label in enumerate(["구분", "분류", "계정명", "연도", "금액"], 1):
        c = data.cell(row=1, column=i, value=label)
        c.font, c.fill = _font(bold=True, color="FFFFFF"), HEAD_FILL
    r = 2
    for ln in ctx.db.bs:
        if ln.cls == m.SUBTOTAL:
            continue
        for year in ctx.years:
            for col, value in enumerate([ln.section, ln.cls, display_name(ln), f"FY{year}", ln.value(year)], 1):
                data.cell(row=r, column=col, value=value).font = _font()
            data.cell(row=r, column=5).number_format = MIL
            r += 1
    last = max(r - 1, 2)
    for col, width in zip("ABCDE", (11, 17, 38, 9, 18)):
        data.column_dimensions[col].width = width

    classes = list(dict.fromkeys(ln.cls for ln in ctx.db.bs if ln.cls != m.SUBTOTAL))
    brief = [
        "DATA 시트는 재무상태표를 '한 줄에 한 값'의 긴 형태로 풀어 놓은 것입니다. 이것으로 피벗 테이블을 만드세요.",
        "① DATA 시트 아무 칸이나 선택 › 삽입 › 피벗 테이블 › 새 워크시트",
        "② 행에 '분류', 열에 '연도', 값에 '금액'(합계)   ③ 아래 검산표를 SUMIFS로 채워 피벗 결과와 같은지 비교",
    ]
    p = Page(ctx, sheet, "과제: 피벗 테이블", brief)
    top = p.top
    p.head(top, _year_heads(ctx, "검산표 (분류)"))
    for k, cls in enumerate(classes):
        row = top + 1 + k
        p.put(f"A{row}", cls)
        for i in range(ctx.n):
            c = L(2 + i)
            p.ans(f"{c}{row}",
                  f"SUMIFS(DATA!$E$2:$E${last},DATA!$B$2:$B${last},$A{row},DATA!$D$2:$D${last},{c}${top})", MIL)
    p.notes(top + len(classes) + 2, [
        "피벗 결과가 이 검산표와 같아야 합니다. 피벗은 탐색용, SUMIFS는 보고서에 남길 표에 씁니다.",
        "피벗은 원본이 바뀌어도 자동으로 바뀌지 않습니다. 피벗 안에서 마우스 오른쪽 › 새로 고침을 눌러야 합니다.",
        "값 칸을 더블클릭하면 그 숫자를 이루는 원본 행들이 새 시트로 펼쳐집니다. 금액의 구성을 볼 때 유용합니다.",
    ])
    return Task("pivot", sheet, "분류 × 연도 피벗과 검산", "피벗 테이블, SUMIFS", brief)


def task_chart(ctx: Ctx, sheet: str) -> Task:
    rev, ebit = ctx.tagged("IS", m.REV), ctx.tagged("IS", m.EBIT)
    brief = [
        "아래 표로 묶은 세로 막대형 차트를 만드세요. (표 선택 › 삽입 › 차트)",
        "① 계열은 매출액·영업이익, 가로축은 연도   ② 차트 제목에 회사명과 단위(억원)를 넣기",
        "③ 눈금선을 옅게, 범례는 아래쪽으로. 보고서에 넣는다고 생각하고 다듬어 보세요.",
    ]
    p = Page(ctx, sheet, "과제: 차트", brief)
    top = p.top
    p.head(top, ["연도", "매출액 (억원)", "영업이익 (억원)"])
    for i, year in enumerate(ctx.years):
        r = top + 1 + i
        p.put(f"A{r}", f"FY{year}")
        p.put(f"B{r}", f"=ROUND({ctx.cell('IS', rev, i)}/100000000,0)" if rev else 0, INT)
        p.put(f"C{r}", f"=ROUND({ctx.cell('IS', ebit, i)}/100000000,0)" if ebit else 0, INT)
    if ctx.answer:
        chart = BarChart()
        chart.type, chart.grouping = "col", "clustered"
        chart.title = f"{ctx.db.corp['corp_name']} 매출액·영업이익 (억원)"
        chart.add_data(Reference(p.ws, min_col=2, max_col=3, min_row=top, max_row=top + ctx.n), titles_from_data=True)
        chart.set_categories(Reference(p.ws, min_col=1, min_row=top + 1, max_row=top + ctx.n))
        chart.legend.position = "b"
        chart.height, chart.width = 8, 16
        p.ws.add_chart(chart, f"E{top}")
    p.notes(top + ctx.n + 2, [
        "머리글까지 포함해 표를 선택하고 차트를 넣으면 계열 이름과 축이 자동으로 잡힙니다.",
        "매출액과 영업이익의 크기 차이가 커서 영업이익 막대가 안 보이면, 두 지표를 한 차트에 넣지 말고 차트를 둘로 나누는 편이 읽기 쉽습니다.",
    ])
    return Task("chart", sheet, "매출액·영업이익 차트 만들기", "차트", brief)


SHORTCUT_MISSIONS = [
    ("IS 시트 C4에서 계정명의 맨 아래 행까지 한 번에 이동", "Ctrl + ↓", "⌘ ↓"),
    ("C4부터 맨 아래·맨 오른쪽까지 한 번에 선택", "Ctrl + Shift + ↓, →", "⌘ ⇧ ↓, →"),
    ("FY 금액 범위를 복사해 새 시트에 값만 붙여넣기", "Ctrl + C, Alt → E → S → V", "⌘ C, ⌃ ⌘ V 후 V"),
    ("붙여넣은 숫자에 천 단위 구분 기호 적용", "Ctrl + Shift + 1", "⌃ ⇧ 1"),
    ("열 너비를 내용에 맞게 자동 조정", "Alt → H → O → I", "열 경계선 더블클릭"),
    ("머리글 행에 필터 걸기", "Ctrl + Shift + L", "⌘ ⇧ F"),
    ("열 하나를 통째로 선택해 왼쪽에 새 열 삽입", "Ctrl + Space, Ctrl + Shift + +", "⌃ Space, ⌃ ⇧ ="),
    ("새 열 맨 위 칸에 수식을 쓰고 아래로 채우기", "범위 선택 후 Ctrl + D", "범위 선택 후 ⌃ D"),
    ("수식의 참조를 절대참조로 바꾸기", "수식 편집 중 F4", "⌘ T 또는 F4"),
    ("행 3개를 그룹으로 묶었다가 풀기", "Alt + Shift + → / ←", "⌘ ⇧ K / ⌘ ⇧ J"),
    ("표 바깥에 테두리 치기", "Ctrl + Shift + 7", "⌘ ⌥ 0"),
    ("1행 아래로 틀 고정", "Alt → W → F → F", "보기 탭 › 틀 고정"),
    ("BS · IS · CF 시트를 차례로 넘겨 보기", "Ctrl + Page Down / Up", "⌥ → / ←"),
]


def task_shortcuts(ctx: Ctx, sheet: str) -> Task:
    brief = [
        "마우스를 쓰지 않고 아래 미션을 순서대로 해 보세요. 연습용이니 이 파일의 사본에서 하면 됩니다.",
        "처음에는 느려도 됩니다. 한 바퀴 돌고 나서 정답 파일의 키와 비교한 뒤, 안 보고 한 번 더 해 보세요.",
    ]
    p = Page(ctx, sheet, "과제: 단축키 미션", brief)
    top = p.top
    p.head(top, ["미션", "Windows", "Mac"])
    p.ws.column_dimensions["A"].width = 52
    p.ws.column_dimensions["B"].width = 34
    p.ws.column_dimensions["C"].width = 30
    for k, (mission, win, mac) in enumerate(SHORTCUT_MISSIONS):
        r = top + 1 + k
        p.put(f"A{r}", f"{k + 1}. {mission}")
        for col, keys in (("B", win), ("C", mac)):
            c = p.ws[f"{col}{r}"]
            c.fill = ANSWER_FILL
            if ctx.answer:
                c.value, c.font = keys, _font(color=GREEN_DARK)
    p.note(top + len(SHORTCUT_MISSIONS) + 2, "'Alt → H → O → I'는 키를 하나씩 순서대로, 'Ctrl + D'는 동시에 누릅니다. Alt 순차 키는 Windows에만 있습니다.")
    return Task("shortcuts", sheet, "마우스 없이 표 다루기 13가지", "단축키", brief)


MACRO_FORMAT = '''Sub FormatReport()
    ' 지금 보고 있는 시트를 보고서 서식으로 정리한다
    Dim ws As Worksheet
    Dim lastRow As Long, lastCol As Long

    Set ws = ActiveSheet
    lastRow = ws.Cells(ws.Rows.Count, 3).End(xlUp).Row          ' C열(계정명)의 마지막 행
    lastCol = ws.Cells(3, ws.Columns.Count).End(xlToLeft).Column ' 3행(머리글)의 마지막 열

    ' 1) 머리글: 굵게 + 연한 초록 채우기
    With ws.Range(ws.Cells(3, 1), ws.Cells(3, lastCol))
        .Font.Bold = True
        .Interior.Color = RGB(226, 239, 218)
    End With

    ' 2) 금액: 백만원 단위, 음수는 괄호
    ws.Range(ws.Cells(4, 4), ws.Cells(lastRow, lastCol)).NumberFormat = "#,##0,,;(#,##0,,);""-"""

    ' 3) 열 너비 자동 맞춤
    ws.Range(ws.Columns(1), ws.Columns(lastCol)).AutoFit

    ' 4) 머리글 아래, 계정명 오른쪽에서 틀 고정
    ws.Cells(4, 4).Select
    ActiveWindow.FreezePanes = False
    ActiveWindow.FreezePanes = True
End Sub
'''

MACRO_CHECK = '''Sub CheckBalance()
    ' BS 시트에서 연도별로 자산총계 = 부채총계 + 자본총계인지 확인해 새 시트에 적는다
    Dim ws As Worksheet, out As Worksheet
    Dim lastRow As Long, lastCol As Long, r As Long, c As Long
    Dim rAsset As Long, rLiab As Long, rEquity As Long
    Dim diff As Double

    Set ws = ThisWorkbook.Worksheets("BS")
    lastRow = ws.Cells(ws.Rows.Count, 3).End(xlUp).Row
    lastCol = ws.Cells(3, ws.Columns.Count).End(xlToLeft).Column

    ' 1) 총계가 몇 행에 있는지 찾는다
    For r = 4 To lastRow
        Select Case ws.Cells(r, 3).Value
            Case "자산총계": rAsset = r
            Case "부채총계": rLiab = r
            Case "자본총계": rEquity = r
        End Select
    Next r
    If rAsset = 0 Or rLiab = 0 Or rEquity = 0 Then
        MsgBox "총계 행을 찾지 못했습니다."
        Exit Sub
    End If

    ' 2) 결과를 적을 시트를 맨 뒤에 만든다
    Set out = ThisWorkbook.Worksheets.Add(After:=ThisWorkbook.Worksheets(ThisWorkbook.Worksheets.Count))
    out.Name = "검증결과"
    out.Range("A1:C1").Value = Array("연도", "차이", "판정")

    ' 3) 연도 열(D열부터)을 하나씩 돌며 차이를 계산한다
    For c = 4 To lastCol
        diff = ws.Cells(rAsset, c).Value - ws.Cells(rLiab, c).Value - ws.Cells(rEquity, c).Value
        out.Cells(c - 2, 1).Value = ws.Cells(3, c).Value
        out.Cells(c - 2, 2).Value = diff
        If Abs(diff) <= 1 Then
            out.Cells(c - 2, 3).Value = "OK"
        Else
            out.Cells(c - 2, 3).Value = "확인 필요"
            out.Cells(c - 2, 3).Font.Color = vbRed
        End If
    Next c
End Sub
'''

MACRO_HOWTO = [
    "실행 방법: ① 이 파일을 '다른 이름으로 저장 › Excel 매크로 사용 통합 문서(.xlsm)'로 저장",
    "② VBA 편집기 열기 (Windows: Alt + F11, Mac: ⌥ F11) › 삽입 › 모듈",
    "③ 아래 코드를 복사해 모듈에 붙여넣기 (함께 받은 '매크로 정답 코드.txt'에도 같은 코드가 있습니다)",
    "④ 엑셀로 돌아와 매크로 목록(Alt + F8 / ⌥ F8)에서 실행. 매크로는 실행 취소가 안 되니 먼저 저장하세요.",
]


def _macro_page(ctx: Ctx, sheet: str, title: str, brief: list[str], code: str, notes: list[str]) -> None:
    p = Page(ctx, sheet, title, brief)
    p.ws.column_dimensions["A"].width = 110
    if not ctx.answer:
        return
    r = p.top
    for text in MACRO_HOWTO:
        p.ws.cell(row=r, column=1, value=text).font = _font(color=GREY)
        r += 1
    r += 1
    for line in code.rstrip("\n").split("\n"):
        c = p.ws.cell(row=r, column=1, value=line or None)
        c.font, c.fill = _font(name="Menlo", color=GREEN_DARK), ANSWER_FILL
        r += 1
    p.notes(r + 1, notes)


def _macro_practice_sheet(ctx: Ctx) -> None:
    """서식이 하나도 없는 손익계산서. 서식 매크로를 돌려 볼 대상이다."""
    if "매크로연습" in ctx.wb.sheetnames:
        return
    ws = ctx.wb.create_sheet("매크로연습")
    ws["A1"] = "서식이 없는 손익계산서입니다. 이 시트를 보고 있는 상태에서 FormatReport 매크로를 실행하세요."
    for i, label in enumerate(["구분", "태그", "계정명"] + [f"FY{y}" for y in ctx.years], 1):
        ws.cell(row=HDR, column=i, value=label)
    for k, ln in enumerate(ctx.db.is_):
        ws.cell(row=FIRST + k, column=3, value=display_name(ln))
        for i, year in enumerate(ctx.years):
            ws.cell(row=FIRST + k, column=YEAR_COL + i, value=ln.values.get(year))


def task_macro_format(ctx: Ctx, sheet: str) -> Task:
    brief = [
        "'매크로연습' 시트는 서식이 없는 손익계산서입니다. 이 시트를 보고서 모양으로 정리하는 매크로 FormatReport를 만드세요.",
        "① 3행 머리글을 굵게 + 연한 초록 채우기   ② 금액(D4부터)을 백만원 단위·음수 괄호 서식으로",
        "③ 열 너비 자동 맞춤   ④ D4 기준으로 틀 고정",
        "먼저 '매크로 기록'으로 같은 조작을 녹화해 코드를 본 뒤, 마지막 행·열을 자동으로 찾도록 고쳐 보세요.",
    ]
    _macro_practice_sheet(ctx)
    _macro_page(ctx, sheet, "과제: 매크로 기초 — 서식 자동화", brief, MACRO_FORMAT, [
        "기록한 매크로는 Range(\"D4:H40\")처럼 범위가 고정됩니다. lastRow·lastCol로 바꾸면 행 수가 다른 회사 자료에도 그대로 쓸 수 있습니다.",
        "With … End With는 같은 대상에 여러 속성을 줄 때 반복을 줄여 줍니다.",
        "서식 코드 안의 큰따옴표는 두 번(\"\") 적어야 문자로 인식됩니다.",
    ])
    return Task("macro_format", sheet, "보고서 서식을 한 번에 입히는 매크로", "매크로 기록, Range, With", brief)


def task_macro_loop(ctx: Ctx, sheet: str) -> Task:
    brief = [
        "BS 시트에서 연도별로 '자산총계 − 부채총계 − 자본총계'를 계산해 새 시트에 적는 매크로 CheckBalance를 만드세요.",
        "① For 반복문으로 C열을 훑어 세 총계가 몇 행인지 찾기   ② '검증결과' 시트를 새로 만들기",
        "③ 연도 열(D열부터)을 하나씩 돌며 연도·차이·판정(차이가 1 이하면 OK, 아니면 '확인 필요'를 빨간 글씨로) 적기",
    ]
    _macro_page(ctx, sheet, "과제: 매크로 응용 — 반복문과 조건문", brief, MACRO_CHECK, [
        "Select Case는 한 값을 여러 경우와 비교할 때 If를 여러 번 쓰는 것보다 읽기 쉽습니다.",
        "'검증결과' 시트가 이미 있으면 이름이 겹쳐 오류가 납니다. 다시 실행하려면 그 시트를 먼저 지우세요. 익숙해지면 있을 때 지우고 새로 만드는 코드를 넣어 보세요.",
        "실무에서는 이런 검증을 수식(체크 시트)으로 남기는 것이 원칙입니다. 이 과제는 반복문·조건문 연습용입니다.",
    ])
    return Task("macro_loop", sheet, "총계 검증을 반복문으로 자동화", "For, Select Case, If", brief)


@dataclass
class Skill:
    key: str
    group: str
    label: str
    build: Callable[[Ctx, str], Task]
    macro: str | None = None


SKILLS: list[Skill] = [
    Skill("sum", G_DAILY, "SUM · SUBTOTAL", task_sum),
    Skill("sumifs", G_DAILY, "SUMIFS · COUNTIFS", task_sumifs),
    Skill("lookup", G_DAILY, "VLOOKUP · XLOOKUP", task_lookup),
    Skill("text", G_DAILY, "LEFT · RIGHT · MID · CONCATENATE · COUNTA", task_text),
    Skill("index_match", G_SOMETIMES, "INDEX · MATCH", task_index_match),
    Skill("transpose", G_SOMETIMES, "TRANSPOSE · TOCOL", task_transpose),
    Skill("indirect", G_SOMETIMES, "INDIRECT", task_indirect),
    Skill("abs_ref", G_EXTRA, "절대·혼합 참조", task_abs_ref),
    Skill("if_error", G_EXTRA, "IF · AND · IFERROR", task_if_error),
    Skill("growth", G_EXTRA, "성장률 · CAGR", task_growth),
    Skill("cond_format", G_FEATURE, "조건부 서식", task_cond_format),
    Skill("pivot", G_FEATURE, "피벗 테이블", task_pivot),
    Skill("chart", G_FEATURE, "차트", task_chart),
    Skill("shortcuts", G_SHORTCUT, "단축키 미션", task_shortcuts),
    Skill("macro_format", G_MACRO, "매크로 기초 (서식 자동화)", task_macro_format, MACRO_FORMAT),
    Skill("macro_loop", G_MACRO, "매크로 응용 (반복문·조건문)", task_macro_loop, MACRO_CHECK),
]
SKILL_GROUPS = list(dict.fromkeys(s.group for s in SKILLS))


def _build(db: Databook, skills: list[Skill], answer: bool, seed: int) -> tuple[bytes, list[Task]]:
    wb = Workbook()
    wb.remove(wb.active)
    ctx = Ctx(db, wb, answer, seed)
    _statement_sheet(ctx, "BS", "재무상태표", "분류")
    _statement_sheet(ctx, "IS", "손익계산서" if db.is_source == "IS" else "포괄손익계산서", "태그")
    _statement_sheet(ctx, "CF", "현금흐름표", "태그")
    tasks = [skill.build(ctx, f"과제{n}") for n, skill in enumerate(skills, 1)]
    _guide_sheet(ctx, tasks)
    # 안내 → 과제 → 재무제표 → 보조 시트 순으로 정렬
    order = ["안내"] + [t.sheet for t in tasks] + ["BS", "IS", "CF"]
    wb._sheets = [wb[n] for n in order] + [ws for ws in wb.worksheets if ws.title not in order]
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), tasks


def build_practice(db: Databook, skill_keys: list[str], seed: int = 0) -> Practice:
    """과제 파일과 정답 파일을 같은 구성으로 만든다. seed가 같으면 같은 문제가 나온다."""
    skills = [s for s in SKILLS if s.key in skill_keys]
    if not skills:
        raise ValueError("익히고 싶은 기능을 하나 이상 고르세요.")
    task_xlsx, tasks = _build(db, skills, answer=False, seed=seed)
    answer_xlsx, _ = _build(db, skills, answer=True, seed=seed)
    macros = [s.macro for s in skills if s.macro]
    code = "Option Explicit\n\n" + "\n".join(macros) if macros else None
    return Practice(task_xlsx, answer_xlsx, tasks, code)
