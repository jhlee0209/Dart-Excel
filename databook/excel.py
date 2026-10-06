"""수식이 살아있는 엑셀 데이터북을 만든다.

BS·IS·CF 시트에는 DART 원본 값(파란 글씨)을 넣고, 나머지 시트는 그 값을 참조하는 수식으로만 채운다.
분석 시트는 원본 시트의 '분류'(B열)를 SUMIFS로 집계하므로 엑셀에서 분류를 바꾸면 결과가 따라 바뀐다.
"""
from __future__ import annotations

import io
from datetime import date
from typing import Callable

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from . import mapping as m
from .model import FS_LABEL, Databook, Line

FONT = "맑은 고딕"
BLUE, BLACK, GREY, WHITE = "0000FF", "000000", "7F7F7F", "FFFFFF"
HEADER_FILL = PatternFill("solid", start_color="404040")
SECTION_FILL = PatternFill("solid", start_color="EDEDED")
INPUT_FILL = PatternFill("solid", start_color="FFF2CC")
WARN_FILL = PatternFill("solid", start_color="F8CBAD")
TOP = Border(top=Side(style="thin"))

# 셀 값은 원 단위로 두고 표시형식으로만 백만원을 보여준다.
NUM = '#,##0,,;(#,##0,,);"-"'
WON = '#,##0;(#,##0);"-"'
PCT = '0.0%;-0.0%;"-"'
MULT = '0.0"x";-0.0"x";"-"'
DAYS = '0"일";-0"일";"-"'

RAW_HEADER, RAW_FIRST = 4, 5  # 원본 시트의 머리글 행, 첫 데이터 행
RAW_YEAR_COL = 5  # 원본 시트의 첫 연도 열(E)
AN_YEAR_COL = 2  # 분석 시트의 첫 연도 열(B)

SHEET_ORDER = ["요약", "BS", "IS", "CF", "NWC", "순차입금", "EBITDA", "비율", "체크"]


def _font(color: str = BLACK, bold: bool = False, size: int = 10, italic: bool = False) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def _raw_col(i: int) -> str:
    return get_column_letter(RAW_YEAR_COL + i)


def _col(i: int) -> str:
    return get_column_letter(AN_YEAR_COL + i)


def _title(ws: Worksheet, db: Databook, heading: str, unit: str = "단위: 백만원") -> None:
    ws["A1"] = heading
    ws["A1"].font = _font(bold=True, size=14)
    ws["A2"] = f"{db.corp['corp_name']} | {FS_LABEL[db.fs_div]}재무제표 | {unit}"
    ws["A2"].font = _font(color=GREY)
    ws.sheet_view.showGridLines = False


def _header(ws: Worksheet, row: int, labels: list[str], n_left: int) -> None:
    for c, label in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.font = _font(color=WHITE, bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="left" if c <= n_left else "right")


class RawSheet:
    """DART 원본 값 시트. B열(분류·태그)이 분석 시트의 집계 기준이 된다."""

    def __init__(self, wb: Workbook, db: Databook, name: str, heading: str,
                 lines: list[Line], options: list[str], label: str, is_bs: bool):
        self.name, self.lines = name, lines
        self.last = RAW_FIRST + max(len(lines), 1) - 1
        self.row_of = {id(ln): RAW_FIRST + i for i, ln in enumerate(lines)}
        ws = self.ws = wb.create_sheet(name)
        _title(ws, db, heading, "단위: 백만원 (셀 값은 원 단위, 표시형식만 백만원)")
        ws["A3"] = f"B열 '{label}'을 바꾸면 분석 시트가 따라 바뀝니다. 파란 글씨 = DART 원본 값."
        ws["A3"].font = _font(color=GREY, italic=True)
        _header(ws, RAW_HEADER, ["구분", label, "계정명", "표준계정ID"] + [f"FY{y}" for y in db.years], 4)

        for line in lines:
            r = self.row_of[id(line)]
            is_total = is_bs and line.cls == m.SUBTOTAL
            ws.cell(row=r, column=1, value=line.section or None).font = _font(color=GREY)
            ws.cell(row=r, column=2, value=line.cls or None).font = _font()
            name_cell = ws.cell(row=r, column=3, value=line.name)
            name_cell.font = _font(bold=is_total)
            name_cell.alignment = Alignment(indent=0 if is_total or not is_bs else 1)
            if m.is_standard_id(line.account_id):
                ws.cell(row=r, column=4, value=line.account_id).font = _font(color=GREY, size=8)
            for i, year in enumerate(db.years):
                cell = ws.cell(row=r, column=RAW_YEAR_COL + i, value=line.values.get(year))
                cell.font = _font(color=BLUE, bold=is_total)
                cell.number_format = NUM

        tag_range = f"B{RAW_FIRST}:B{self.last}"
        validation = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        validation.error = "목록에 있는 값만 쓸 수 있습니다."
        ws.add_data_validation(validation)
        validation.add(tag_range)
        ws.conditional_formatting.add(
            tag_range, CellIsRule(operator="equal", formula=[f'"{m.UNCLASSIFIED}"'], fill=WARN_FILL)
        )

        ws.column_dimensions["A"].width = 11
        ws.column_dimensions["A"].hidden = not is_bs
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 42
        ws.column_dimensions["D"].width = 14
        for i in range(len(db.years)):
            ws.column_dimensions[_raw_col(i)].width = 14
        ws.freeze_panes = ws.cell(row=RAW_FIRST, column=RAW_YEAR_COL)

    def _range(self, col: str) -> str:
        return f"'{self.name}'!${col}${RAW_FIRST}:${col}${self.last}"

    def sumifs(self, i: int, tag: str) -> str:
        """i번째 연도 열에서 B열이 tag인 줄의 합."""
        return f'SUMIFS({self._range(_raw_col(i))},{self._range("B")},"{tag}")'

    def section_sum(self, i: int, suffix: str) -> str:
        """합계 줄을 뺀 구분별 합. suffix는 '자산'·'부채'·'자본'."""
        pattern = suffix if suffix == m.SEC_E else f"*{suffix}"
        return (
            f'SUMIFS({self._range(_raw_col(i))},{self._range("A")},"{pattern}",'
            f'{self._range("B")},"<>{m.SUBTOTAL}")'
        )

    def cell(self, line: Line | None, i: int) -> str:
        if line is None:
            return "0"
        return f"'{self.name}'!{_raw_col(i)}{self.row_of[id(line)]}"

    def tag_cell(self, line: Line) -> str:
        return f"'{self.name}'!$B${self.row_of[id(line)]}"


class Sheet:
    """수식만 들어가는 분석 시트. 한 줄씩 아래로 쌓는다."""

    def __init__(self, wb: Workbook, db: Databook, name: str, heading: str,
                 header_row: int = 4, unit: str = "단위: 백만원"):
        self.name, self.db = name, db
        self.ws = wb.create_sheet(name)
        self.rows: dict[str, int] = {}
        self.note_col = AN_YEAR_COL + len(db.years)
        _title(self.ws, db, heading, unit)
        _header(self.ws, header_row, [""] + [f"FY{y}" for y in db.years] + ["비고"], 1)
        self.ws.cell(row=header_row, column=self.note_col).alignment = Alignment(horizontal="left")
        self.row = self.first_row = header_row + 1
        self.ws.column_dimensions["A"].width = 38
        for i in range(len(db.years)):
            self.ws.column_dimensions[_col(i)].width = 14
        self.ws.column_dimensions[get_column_letter(self.note_col)].width = 60
        self.ws.freeze_panes = self.ws.cell(row=header_row + 1, column=AN_YEAR_COL)

    def section(self, label: str) -> None:
        if self.row > self.first_row:
            self.row += 1
        for c in range(1, self.note_col + 1):
            self.ws.cell(row=self.row, column=c).fill = SECTION_FILL
        self.ws.cell(row=self.row, column=1, value=label).font = _font(bold=True)
        self.row += 1

    def line(self, label: str, formula: Callable[[int], str | None], *, key: str | None = None,
             fmt: str = NUM, bold: bool = False, indent: int = 0, note: str | None = None) -> int:
        """formula(i)는 i번째 연도 칸에 넣을 수식('=' 제외). None이면 빈칸."""
        r = self.row
        label_cell = self.ws.cell(row=r, column=1, value=label)
        label_cell.font = _font(bold=bold)
        label_cell.alignment = Alignment(indent=indent)
        for i in range(len(self.db.years)):
            body = formula(i)
            cell = self.ws.cell(row=r, column=AN_YEAR_COL + i, value=f"={body}" if body else None)
            cell.font = _font(bold=bold)
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal="right")
            if bold:
                cell.border = TOP
        if bold:
            label_cell.border = TOP
        if note:
            self.ws.cell(row=r, column=self.note_col, value=note).font = _font(color=GREY, italic=True)
        if key:
            self.rows[key] = r
        self.row += 1
        return r

    def input_line(self, label: str, *, key: str, note: str) -> int:
        """사용자가 직접 채우는 노란 칸(기본값 0)."""
        r = self.row
        self.ws.cell(row=r, column=1, value=label).font = _font(color=BLUE)
        self.ws.cell(row=r, column=1).alignment = Alignment(indent=1)
        for i in range(len(self.db.years)):
            cell = self.ws.cell(row=r, column=AN_YEAR_COL + i, value=0)
            cell.font, cell.fill, cell.number_format = _font(color=BLUE), INPUT_FILL, NUM
        self.ws.cell(row=r, column=self.note_col, value=note).font = _font(color=GREY, italic=True)
        self.rows[key] = r
        self.row += 1
        return r

    def at(self, key: str, i: int) -> str:
        """같은 시트 안에서 쓰는 셀 주소."""
        return f"{_col(i)}{self.rows[key]}"

    def ref(self, key: str, i: int) -> str:
        """다른 시트에서 쓰는 셀 주소."""
        return f"'{self.name}'!{_col(i)}{self.rows[key]}"

    def class_block(self, bs: RawSheet, cls: str, note: str | None = None) -> None:
        """한 분류에 속한 BS 계정을 나열하고 합계를 낸다.

        합계는 항상 SUMIFS라서, 엑셀에서 다른 계정을 이 분류로 옮기면 '기타' 줄에 잡힌다.
        """
        members = [ln for ln in bs.lines if ln.cls == cls]
        first = self.row
        total_row = first + len(members) + (1 if members else 0)
        for ln in members:
            self.line(
                ln.name,
                lambda i, ln=ln: f'IF({bs.tag_cell(ln)}="{cls}",{bs.cell(ln, i)},0)',
                indent=1,
            )
        if members:
            self.line(
                "기타 (엑셀에서 분류를 바꾼 계정)",
                lambda i: f"{_col(i)}{total_row}-SUM({_col(i)}{first}:{_col(i)}{total_row - 2})",
                indent=1,
            )
        self.line(cls if not members else f"{cls} 합계", lambda i: bs.sumifs(i, cls),
                  key=cls, bold=bool(members), note=note)


def _ratio(num: str, den: str) -> str:
    return f'IFERROR(({num})/({den}),"n/a")'


def build_workbook(db: Databook) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    bs = RawSheet(wb, db, "BS", "재무상태표", db.bs, m.BS_CLASSES, "분류", is_bs=True)
    is_title = "손익계산서" if db.is_source == "IS" else "포괄손익계산서"
    is_ = RawSheet(wb, db, "IS", is_title, db.is_, m.IS_TAGS, "태그", is_bs=False)
    cf = RawSheet(wb, db, "CF", "현금흐름표", db.cf, m.CF_TAGS, "태그", is_bs=False)

    # ── EBITDA ──
    eb = Sheet(wb, db, "EBITDA", "EBITDA")
    eb.line("매출액", lambda i: is_.sumifs(i, m.REV), key="매출액")
    eb.line("매출원가", lambda i: is_.sumifs(i, m.COGS), key="매출원가")
    eb.line("매출총이익", lambda i: f"{eb.at('매출액', i)}-{eb.at('매출원가', i)}", key="매출총이익", bold=True)
    eb.line("판매비와관리비", lambda i: is_.sumifs(i, m.SGA), key="판관비")
    eb.line("영업이익 (EBIT)", lambda i: is_.sumifs(i, m.EBIT), key="영업이익", bold=True,
            note="손익계산서에 보고된 영업이익")
    eb.line("감가상각비·상각비", lambda i: cf.sumifs(i, m.DA), key="감가상각비", indent=1,
            note="현금흐름표에서 태그가 '감가상각비'인 줄의 합")
    eb.input_line("감가상각비 직접 입력", key="감가상각비_입력",
                  note="현금흐름표에 없으면 주석의 금액을 원 단위로 입력")
    eb.line("EBITDA", lambda i: f"{eb.at('영업이익', i)}+{eb.at('감가상각비', i)}+{eb.at('감가상각비_입력', i)}",
            key="EBITDA", bold=True)
    eb.line("EBITDA 마진", lambda i: _ratio(eb.at("EBITDA", i), eb.at("매출액", i)), fmt=PCT)
    eb.section("이익의 질(QoE) 조정 연습")
    adj_first = eb.row
    for n in range(1, 4):
        eb.input_line(f"조정 {n} (항목명을 적으세요)", key=f"조정{n}",
                      note="일회성·비경상 항목. 이익을 늘리면 +, 줄이면 − (원 단위)")
    eb.line("조정 후 EBITDA",
            lambda i: f"{eb.at('EBITDA', i)}+SUM({_col(i)}{adj_first}:{_col(i)}{adj_first + 2})",
            key="조정EBITDA", bold=True)
    eb.line("조정 후 EBITDA 마진", lambda i: _ratio(eb.at("조정EBITDA", i), eb.at("매출액", i)), fmt=PCT)
    eb.section("참고")
    eb.line("세전이익", lambda i: is_.sumifs(i, m.PBT), key="세전이익")
    eb.line("법인세비용", lambda i: is_.sumifs(i, m.TAX), key="법인세비용")
    eb.line("당기순이익", lambda i: is_.sumifs(i, m.NI), key="당기순이익")

    # ── NWC ──
    nwc = Sheet(wb, db, "NWC", "순운전자본 (Net Working Capital)")
    nwc.section("운전자본 자산")
    for cls in m.NWC_ASSET_CLASSES:
        nwc.class_block(bs, cls)
    nwc.line("운전자본 자산 합계",
             lambda i: "+".join(nwc.at(c, i) for c in m.NWC_ASSET_CLASSES), key="자산합계", bold=True)
    nwc.section("운전자본 부채")
    for cls in m.NWC_LIAB_CLASSES:
        nwc.class_block(bs, cls)
    nwc.line("운전자본 부채 합계",
             lambda i: "+".join(nwc.at(c, i) for c in m.NWC_LIAB_CLASSES), key="부채합계", bold=True)
    nwc.section("순운전자본")
    nwc.line("순운전자본 (NWC)", lambda i: f"{nwc.at('자산합계', i)}-{nwc.at('부채합계', i)}",
             key="NWC", bold=True)
    nwc.line("전년 대비 증감", lambda i: f"{nwc.at('NWC', i)}-{nwc.at('NWC', i - 1)}" if i else None, indent=1)
    nwc.line("NWC / 매출액", lambda i: _ratio(nwc.at("NWC", i), eb.ref("매출액", i)), fmt=PCT, key="NWC%")
    nwc.section("회전일수")
    nwc.line("매출채권 회전일수 (DSO)", lambda i: _ratio(f"{nwc.at(m.AR, i)}*365", eb.ref("매출액", i)),
             fmt=DAYS, key="DSO", note="매출채권 ÷ 매출액 × 365 (기말 잔액 기준)")
    nwc.line("재고자산 회전일수 (DIO)", lambda i: _ratio(f"{nwc.at(m.INV, i)}*365", eb.ref("매출원가", i)),
             fmt=DAYS, key="DIO", note="재고자산 ÷ 매출원가 × 365")
    nwc.line("매입채무 회전일수 (DPO)", lambda i: _ratio(f"{nwc.at(m.AP, i)}*365", eb.ref("매출원가", i)),
             fmt=DAYS, key="DPO", note="매입채무 ÷ 매출원가 × 365")
    nwc.line("현금전환주기 (CCC)",
             lambda i: f'IFERROR({nwc.at("DSO", i)}+{nwc.at("DIO", i)}-{nwc.at("DPO", i)},"n/a")',
             fmt=DAYS, bold=True, note="DSO + DIO − DPO")

    # ── 순차입금 ──
    nd = Sheet(wb, db, "순차입금", "순차입금 (Net Debt)")
    nd.section("차입금")
    nd.class_block(bs, m.DEBT)
    nd.section("현금성자산")
    nd.class_block(bs, m.CASH)
    nd.section("순차입금")
    nd.line("순차입금", lambda i: f"{nd.at(m.DEBT, i)}-{nd.at(m.CASH, i)}", key="순차입금", bold=True,
            note="음수면 순현금")
    nd.class_block(bs, m.DEBT_LIKE,
                   note="퇴직급여부채·미지급법인세 등. BS 시트에서 직접 '유사차입금'으로 분류해 보세요.")
    nd.line("순차입금 (유사차입금 포함)", lambda i: f"{nd.at('순차입금', i)}+{nd.at(m.DEBT_LIKE, i)}",
            key="순차입금_조정", bold=True)
    nd.section("배수")
    nd.line("순차입금 / EBITDA", lambda i: _ratio(nd.at("순차입금", i), eb.ref("EBITDA", i)), fmt=MULT,
            key="배수")
    nd.line("순차입금 / 자본총계",
            lambda i: _ratio(nd.at("순차입금", i), bs.cell(db.role_line("자본총계"), i)), fmt=PCT)

    # ── 비율 ──
    def total(role: str, i: int) -> str:
        return bs.cell(db.role_line(role), i)

    capex = lambda i: f"ABS({cf.sumifs(i, m.CAPEX_PPE)})+ABS({cf.sumifs(i, m.CAPEX_INT)})"  # noqa: E731
    rt = Sheet(wb, db, "비율", "주요 비율", unit="금액은 백만원")
    rt.section("성장성")
    rt.line("매출 성장률",
            lambda i: _ratio(f"{eb.ref('매출액', i)}-{eb.ref('매출액', i - 1)}", eb.ref("매출액", i - 1)) if i else None,
            fmt=PCT, key="성장률")
    rt.line("영업이익 성장률",
            lambda i: _ratio(f"{eb.ref('영업이익', i)}-{eb.ref('영업이익', i - 1)}", f"ABS({eb.ref('영업이익', i - 1)})") if i else None,
            fmt=PCT)
    rt.section("수익성")
    rt.line("매출총이익률", lambda i: _ratio(eb.ref("매출총이익", i), eb.ref("매출액", i)), fmt=PCT)
    rt.line("영업이익률", lambda i: _ratio(eb.ref("영업이익", i), eb.ref("매출액", i)), fmt=PCT, key="영업이익률")
    rt.line("EBITDA 마진", lambda i: _ratio(eb.ref("EBITDA", i), eb.ref("매출액", i)), fmt=PCT, key="EBITDA마진")
    rt.line("순이익률", lambda i: _ratio(eb.ref("당기순이익", i), eb.ref("매출액", i)), fmt=PCT)
    rt.line("ROE", lambda i: _ratio(eb.ref("당기순이익", i), total("자본총계", i)), fmt=PCT,
            note="당기순이익 ÷ 기말 자본총계")
    rt.section("안정성")
    rt.line("부채비율", lambda i: _ratio(total("부채총계", i), total("자본총계", i)), fmt=PCT)
    rt.line("유동비율", lambda i: _ratio(total(m.CUR_A, i), total(m.CUR_L, i)), fmt=PCT)
    rt.line("순차입금 / EBITDA", lambda i: nd.ref("배수", i), fmt=MULT)
    rt.section("현금흐름 (단위: 백만원)")
    rt.line("영업활동현금흐름", lambda i: cf.sumifs(i, m.CFO), key="CFO")
    rt.line("CAPEX (유·무형자산 취득)", capex, key="CAPEX")
    rt.line("FCF", lambda i: f"{rt.at('CFO', i)}-{rt.at('CAPEX', i)}", key="FCF", bold=True,
            note="영업활동현금흐름 − CAPEX (간이 계산)")
    rt.line("CAPEX / 매출액", lambda i: _ratio(rt.at("CAPEX", i), eb.ref("매출액", i)), fmt=PCT)
    rt.line("영업활동현금흐름 / EBITDA", lambda i: _ratio(rt.at("CFO", i), eb.ref("EBITDA", i)), fmt=PCT,
            note="현금전환율")

    # ── 체크 ──
    ck = Sheet(wb, db, "체크", "검증", unit="단위: 원 (0이어야 정상)")
    ck.section("재무상태표")
    ck.line("자산총계 − 부채총계 − 자본총계",
            lambda i: f"{total('자산총계', i)}-{total('부채총계', i)}-{total('자본총계', i)}", fmt=WON, key="c1")
    ck.line("자산 계정 합계 − 자산총계", lambda i: f"{bs.section_sum(i, '자산')}-{total('자산총계', i)}",
            fmt=WON, key="c2", note="0이 아니면 BS에 '소계'로 표시해야 할 줄이 더 있거나 빠진 줄이 있음")
    ck.line("부채 계정 합계 − 부채총계", lambda i: f"{bs.section_sum(i, '부채')}-{total('부채총계', i)}",
            fmt=WON, key="c3")
    ck.line("자본 계정 합계 − 자본총계", lambda i: f"{bs.section_sum(i, '자본')}-{total('자본총계', i)}",
            fmt=WON, key="c4")
    checks = ["c1", "c2", "c3", "c4"]
    ck.line("판정",
            lambda i: "IF(AND(" + ",".join(f"ABS({ck.at(k, i)})<=1" for k in checks) + '),"OK","확인 필요")',
            key="판정", bold=True)
    ck.section("참고")
    ck.line("미분류 금액 (자산·부채)", lambda i: bs.sumifs(i, m.UNCLASSIFIED), fmt=WON,
            note="BS 시트에서 주황색으로 표시된 계정. 성격을 판단해 분류를 정하세요.")
    ck.line("매출액 − 매출원가 − 판관비 − 영업이익",
            lambda i: f"{eb.ref('매출액', i)}-{eb.ref('매출원가', i)}-{eb.ref('판관비', i)}-{eb.ref('영업이익', i)}",
            fmt=WON, note="0이 아니면 기타영업손익 등 다른 영업 항목이 있다는 뜻 (오류는 아님)")

    # ── 요약 ──
    sm = Sheet(wb, db, "요약", f"{db.corp['corp_name']} 미니 데이터북", header_row=8)
    info = [
        ("종목코드", db.corp.get("stock_code") or "비상장"),
        ("재무제표", f"{FS_LABEL[db.fs_div]} / 사업보고서 기준 FY{db.years[0]}~FY{db.years[-1]}"),
        ("출처", "금융감독원 전자공시시스템(OpenDART), 각 연도 사업보고서의 당기 금액"),
        ("생성일", date.today().isoformat()),
    ]
    for r, (label, value) in enumerate(info, 3):
        sm.ws.cell(row=r, column=1, value=label).font = _font(color=GREY)
        sm.ws.cell(row=r, column=2, value=value).font = _font()
    sm.line("매출액", lambda i: eb.ref("매출액", i))
    sm.line("매출 성장률", lambda i: rt.ref("성장률", i) if i else None, fmt=PCT, indent=1)
    sm.line("영업이익", lambda i: eb.ref("영업이익", i))
    sm.line("영업이익률", lambda i: rt.ref("영업이익률", i), fmt=PCT, indent=1)
    sm.line("EBITDA", lambda i: eb.ref("EBITDA", i), bold=True)
    sm.line("EBITDA 마진", lambda i: rt.ref("EBITDA마진", i), fmt=PCT, indent=1)
    sm.line("당기순이익", lambda i: eb.ref("당기순이익", i))
    sm.row += 1
    sm.line("순운전자본 (NWC)", lambda i: nwc.ref("NWC", i), bold=True)
    sm.line("NWC / 매출액", lambda i: nwc.ref("NWC%", i), fmt=PCT, indent=1)
    sm.line("순차입금", lambda i: nd.ref("순차입금", i), bold=True, note="음수면 순현금")
    sm.line("순차입금 / EBITDA", lambda i: nd.ref("배수", i), fmt=MULT, indent=1)
    sm.line("FCF", lambda i: rt.ref("FCF", i))
    sm.row += 1
    sm.line("검증", lambda i: ck.ref("판정", i), note="'확인 필요'면 체크 시트를 보세요")

    notes = [
        "읽는 법",
        "· 파란 글씨는 DART 원본 값, 검은 글씨는 수식, 노란 칸은 직접 입력하는 칸입니다.",
        "· BS 시트 B열의 분류를 바꾸면 NWC·순차입금이 따라 바뀝니다. 주황색 '미분류' 계정부터 판단해 보세요.",
        "· 과거 연도는 그 해 사업보고서에 처음 보고된 금액이라, 이후 재작성된 금액과 다를 수 있습니다.",
        "· 자동 분류는 계정명 기준의 출발점일 뿐입니다. 주석을 읽고 성격에 맞게 고치는 것이 실사의 핵심입니다.",
    ] + [f"· 주의: {w}" for w in db.warnings]
    for offset, text in enumerate(notes):
        cell = sm.ws.cell(row=sm.row + 1 + offset, column=1, value=text)
        cell.font = _font(bold=offset == 0, color=BLACK if offset == 0 else GREY)

    wb._sheets = [wb[name] for name in SHEET_ORDER]
    wb.active = 0

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
